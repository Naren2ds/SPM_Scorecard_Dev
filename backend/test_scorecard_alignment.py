"""
test_scorecard_alignment.py
============================
Verifies that per-KPI earned scores in the Normalized Scorecard match the
scores computed independently from the raw KPI data for N random parent
suppliers.

Two modes
---------
  --offline   Load CSV files directly from backend/data/ (no server needed).
              Uses the same scorecard.py module the server uses.

  (default)   Call the live API on http://127.0.0.1:8000, fetch raw KPI rows,
              aggregate them independently, and compare.

Usage
-----
    # With server running (default):
    python backend/test_scorecard_alignment.py

    # Without server (CSV-only):
    python backend/test_scorecard_alignment.py --offline

    # 50 suppliers, save to Excel:
    python backend/test_scorecard_alignment.py --offline --n 50 --excel
    python backend/test_scorecard_alignment.py --n 50 --excel

    # Custom options:
    python backend/test_scorecard_alignment.py --n 50 --seed 99
    python backend/test_scorecard_alignment.py --api http://127.0.0.1:8000 --n 20
"""

from __future__ import annotations

import argparse
import math
import random
import sys
from pathlib import Path

import pandas as pd

# ── Path setup ─────────────────────────────────────────────────────────────
SCRIPT_DIR = Path(__file__).resolve().parent
DATA_DIR = SCRIPT_DIR / "data"
sys.path.insert(0, str(SCRIPT_DIR))

# ── Tolerance ──────────────────────────────────────────────────────────────
# Earned scores may differ by a tiny floating-point rounding amount.
TOLERANCE = 0.005  # 0.005 points out of 10 max = 0.05% tolerance

# ── KPI definitions (mirrors scorecard.py KPI_CONFIGS exactly) ─────────────
KPI_DEFS: list[dict] = [
    {"id": "DOT",  "name": "Delivery On Time",        "max": 10.0, "floor": 0.70, "target": 0.85, "direction": "higher",   "csv": "dot_kpi.csv"},
    {"id": "SA",   "name": "Supplier Assessment",      "max": 10.0, "floor": 0.50, "target": 0.80, "direction": "higher",   "csv": "supplier_assessment.csv"},
    {"id": "SC",   "name": "Supplier Compliance",      "max": 5.0,  "floor": 0.60, "target": 0.90, "direction": "higher",   "csv": "supplier_compliance.csv"},
    {"id": "PDIV", "name": "Price Divergence",         "max": 5.0,  "floor": 0.05, "target": 0.02, "direction": "lower",    "csv": "price_divergence.csv"},
    {"id": "IC",   "name": "Invoice Conformity",       "max": 5.0,  "floor": 0.80, "target": 0.95, "direction": "higher",   "csv": "invoice_conformity.csv"},
    {"id": "IOT",  "name": "Invoice On Time",          "max": 10.0, "floor": 0.70, "target": 0.85, "direction": "higher",   "csv": "iot_kpi.csv"},
    {"id": "SM",   "name": "Supplier Maturity",        "max": 10.0, "floor": 0.40, "target": 0.80, "direction": "higher",   "csv": "supplier_maturity.csv"},
    {"id": "ECL",  "name": "Eclipse Score",            "max": 5.0,  "floor": 0.50, "target": 0.80, "direction": "higher",   "csv": "eclipse.csv"},
    {"id": "CO2",  "name": "CO2 Reduction Potential",  "max": 5.0,  "floor": None, "target": None, "direction": "quartile", "csv": "co2_emission.csv"},
]

# ── Pure helper functions (independent of scorecard.py) ────────────────────

def _to_num(v) -> float | None:
    """Best-effort numeric coercion, returns None on failure."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        f = float(v)
        return None if math.isnan(f) else f
    s = str(v).strip().replace(",", "").replace("%", "")
    if not s or s.lower() in {"nan", "none", "null", ""}:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _parent_key(row: dict) -> str:
    parent = str(row.get("parentSupplier", "") or "").strip()
    if parent and parent.lower() != "none":
        return parent
    return str(row.get("supplier", "") or "").strip() or "(Unknown)"


def _is_applicable(v) -> bool:
    return str(v or "Applicable").strip().lower() != "not applicable"


def _attainment(raw: float, floor: float, target: float, direction: str) -> float:
    if direction == "higher":
        if raw >= target:
            return 1.0
        if raw <= floor:
            return 0.0
        return (raw - floor) / (target - floor)
    else:  # lower
        if raw <= target:
            return 1.0
        if raw >= floor:
            return 0.0
        return (floor - raw) / (floor - target)


def _aggregate(kid: str, rows: list[dict]) -> dict[str, float]:
    """
    Aggregate raw rows into a per-parent ratio dict.
    Replicates the exact logic in scorecard.py _aggregate_kpi().
    """
    agg: dict[str, dict] = {}
    for r in rows:
        if not _is_applicable(r.get("kpiApplicability")):
            continue
        key = _parent_key(r)
        b = agg.setdefault(key, {"num": 0.0, "den": 0.0, "sum": 0.0, "n": 0.0})

        if kid == "DOT":
            on_time  = _to_num(r.get("onTimePoLines"))            or 0.0
            total    = _to_num(r.get("totalDeliveredPoLines"))    or 0.0
            delayed  = _to_num(r.get("x1DelayedOver30Days"))      or 0.0
            early    = _to_num(r.get("x2EarlyOver30Days"))        or 0.0
            b["num"] += on_time
            b["den"] += total + 0.99 * delayed + 0.10 * early

        elif kid == "IOT":
            b["num"] += _to_num(r.get("invoiceOnTimeCount")) or 0.0
            b["den"] += _to_num(r.get("totalPoLines"))       or 0.0

        elif kid == "IC":
            total_inv  = _to_num(r.get("totalInvoices"))   or 0.0
            mismatches = _to_num(r.get("mismatchCount"))   or 0.0
            b["num"] += max(total_inv - mismatches, 0.0)
            b["den"] += total_inv

        elif kid == "PDIV":
            po  = _to_num(r.get("poValue"))
            inv = _to_num(r.get("invoiceValue"))
            if po is None or inv is None or po <= 0:
                continue
            b["num"] += abs(inv - po)
            b["den"] += po

        elif kid == "SA":
            g   = _to_num(r.get("greenCount"))  or 0.0
            y   = _to_num(r.get("yellowCount")) or 0.0
            red = _to_num(r.get("redCount"))    or 0.0
            valid = g + y + red
            if valid <= 0:
                continue
            b["num"] += g * 1.0 + y * 0.5
            b["den"] += valid

        elif kid == "SC":
            v = _to_num(r.get("compliancePct"))
            if v is None:
                continue
            b["sum"] += v;  b["n"] += 1.0

        elif kid == "SM":
            v = _to_num(r.get("maturityScore"))
            if v is None:
                continue
            if v > 1.0:
                v /= 100.0
            b["sum"] += v;  b["n"] += 1.0

        elif kid == "ECL":
            v = _to_num(r.get("eclipseScore"))
            if v is None:
                continue
            if v > 1.0:
                v /= 100.0
            b["sum"] += v;  b["n"] += 1.0

        elif kid == "CO2":
            v = _to_num(r.get("co2Emission"))
            if v is None:
                continue
            b["sum"] += v;  b["n"] += 1.0

    ratios: dict[str, float] = {}
    for key, b in agg.items():
        if kid in {"SC", "SM", "ECL", "CO2"}:
            if b["n"] > 0:
                ratios[key] = b["sum"] / b["n"]
        else:
            if b["den"] > 0:
                ratios[key] = b["num"] / b["den"]
    return ratios


def _softstretch_earned(
    ratios: dict[str, float],
    direction: str,
    floor,
    target,
    max_score: float,
) -> dict[str, float]:
    """
    Apply softStretch formula to ratios → earned scores.
    Matches scorecard.py _kpi_attainments() exactly.
      earned = max_score × attainment × (0.70 + 0.30 × percentile)
    """
    if direction == "quartile":
        values = list(ratios.values())
        if len(values) >= 2:
            s = pd.Series(values, dtype=float)
            q1, q3 = float(s.quantile(0.25)), float(s.quantile(0.75))
            if q3 <= q1:
                q3 = q1 + 1e-9
            floor, target = q1, q3
        else:
            floor, target = 0.0, 1.0
        direction = "higher"

    attainments = {
        k: _attainment(v, float(floor), float(target), direction)
        for k, v in ratios.items()
    }

    # Best performer → highest percentile (rank 0-indexed from best)
    reverse_sort = direction != "lower"
    sorted_keys = sorted(ratios, key=lambda k: ratios[k], reverse=reverse_sort)
    total = len(sorted_keys)
    percentiles = {k: (total - i) / total for i, k in enumerate(sorted_keys)}

    return {
        k: max_score * attainments[k] * (0.70 + 0.30 * percentiles[k])
        for k in ratios
    }


# ── Test runner ────────────────────────────────────────────────────────────

def _print_header(n: int, mode: str) -> None:
    print(f"\n{'=' * 110}")
    print(f"  Scorecard ↔ Individual KPI Alignment Test  |  Mode: {mode}  |  Sample: {n} random parent suppliers")
    print(f"  Tolerance: ±{TOLERANCE} earned points  |  Formula: softStretch (attainment × 70%+30%×percentile)")
    print(f"{'=' * 110}\n")
    print(f"{'Parent Supplier':<42} {'KPI':<6} {'KPI Name':<28} {'Scorecard':>10} {'IndepCalc':>10} {'Delta':>8}  Status")
    print("-" * 110)


def _compare_and_print(
    sample: list[str],
    sc_lookup: dict[str, dict[str, float]],  # parent → {kpi_id → earned}
    ind_earned: dict[str, dict[str, float]],  # kpi_id  → {parent → earned}
) -> tuple[int, int, list[dict]]:
    """Print comparison rows. Returns (passed, failed, detail_rows)."""
    passed = failed = 0
    detail_rows: list[dict] = []

    for parent in sorted(sample):
        sc_kpis = sc_lookup.get(parent, {})
        any_row = False

        for kpi in KPI_DEFS:
            kid = kpi["id"]
            sc_val  = sc_kpis.get(kid)
            ind_val = ind_earned.get(kid, {}).get(parent)

            # Skip if neither side has data for this KPI
            if sc_val is None and ind_val is None:
                continue

            any_row = True

            sc_str  = f"{sc_val:.4f}"  if sc_val  is not None else "   —   "
            ind_str = f"{ind_val:.4f}" if ind_val is not None else "   —   "

            if sc_val is None:
                delta_str = "  N/A  "
                status    = "MISSING_SCORECARD"
                failed += 1
            elif ind_val is None:
                delta_str = "  N/A  "
                status    = "MISSING_INDEP"
                failed += 1
            else:
                delta = abs(sc_val - ind_val)
                delta_str = f"{delta:>8.4f}"
                if delta <= TOLERANCE:
                    status = "PASS"
                    passed += 1
                else:
                    status = f"FAIL  (diff={delta:.4f})"
                    failed += 1

            print(
                f"{parent:<42} {kid:<6} {kpi['name']:<28} "
                f"{sc_str:>10} {ind_str:>10} {delta_str}  {status}"
            )

            detail_rows.append({
                "Parent Supplier": parent,
                "KPI ID": kid,
                "KPI Name": kpi["name"],
                "Max Score": kpi["max"],
                "Scorecard Earned": round(sc_val, 6) if sc_val is not None else None,
                "Independent Calc Earned": round(ind_val, 6) if ind_val is not None else None,
                "Delta": round(abs(sc_val - ind_val), 6) if (sc_val is not None and ind_val is not None) else None,
                "Tolerance": TOLERANCE,
                "Status": "PASS" if status == "PASS" else status,
            })

        if any_row:
            print()  # blank line between parents

    return passed, failed, detail_rows


def _print_summary(passed: int, failed: int) -> None:
    total = passed + failed
    pct   = 100.0 * passed / total if total else 0.0
    print("-" * 110)
    print(f"\n  RESULT: {passed}/{total} checks passed ({pct:.1f}%)")
    if failed:
        print(f"  {failed} check(s) FAILED — scores differ by more than ±{TOLERANCE}")
        print("  Typical causes:")
        print("    • Year filter mismatch (individual KPI pages default to 2025+2026; scorecard uses all years)")
        print("    • Month filter difference (DOT/IOT have month selectors)")
        print("    • KPI page uses a non-default config (floor/target/maxScore changed by user)")
    else:
        print("  All scores match within tolerance. ✓")
    print()


# ── Excel export ────────────────────────────────────────────────────────────

def _export_excel(
    detail_rows: list[dict],
    sample: list[str],
    passed: int,
    failed: int,
    mode: str,
    seed: int,
    out_path: Path,
) -> None:
    """Write a two-sheet Excel report: Detail + Summary."""
    try:
        import openpyxl  # type: ignore[import]
        from openpyxl.styles import (
            Alignment, Border, Font, PatternFill, Side
        )
        from openpyxl.utils import get_column_letter
        from datetime import datetime
    except ImportError:
        print("  [WARNING] openpyxl not installed — skipping Excel export.")
        print("            Run: pip install openpyxl")
        return

    wb = openpyxl.Workbook()

    # ── Colour palette ──────────────────────────────────────────────────────
    GREEN_FILL  = PatternFill("solid", fgColor="C6EFCE")
    RED_FILL    = PatternFill("solid", fgColor="FFC7CE")
    AMBER_FILL  = PatternFill("solid", fgColor="FFEB9C")
    HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
    ALT_FILL    = PatternFill("solid", fgColor="EEF3FA")

    HEADER_FONT  = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    BODY_FONT    = Font(name="Calibri", size=10)
    BOLD_FONT    = Font(name="Calibri", bold=True, size=10)
    TITLE_FONT   = Font(name="Calibri", bold=True, size=13, color="1F4E79")

    THIN = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    CENTER = Alignment(horizontal="center", vertical="center")
    LEFT   = Alignment(horizontal="left",   vertical="center")

    def hdr(ws, row, col, value):
        c = ws.cell(row=row, column=col, value=value)
        c.fill   = HEADER_FILL
        c.font   = HEADER_FONT
        c.border = BORDER
        c.alignment = CENTER
        return c

    def cell(ws, row, col, value, fill=None, bold=False, align=None):
        c = ws.cell(row=row, column=col, value=value)
        c.font   = BOLD_FONT if bold else BODY_FONT
        c.border = BORDER
        c.alignment = align or CENTER
        if fill:
            c.fill = fill
        return c

    # ════════════════════════════════════════════════════════════════════════
    # Sheet 1 — Detail
    # ════════════════════════════════════════════════════════════════════════
    ws1 = wb.active
    ws1.title = "Detail"
    ws1.freeze_panes = "A3"

    # Title row
    ws1.merge_cells("A1:I1")
    title_cell = ws1["A1"]
    title_cell.value = (
        f"Scorecard ↔ Individual KPI Alignment Report  |  "
        f"{len(sample)} Parent Suppliers  |  Mode: {mode}  |  "
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    title_cell.font = TITLE_FONT
    title_cell.alignment = LEFT
    ws1.row_dimensions[1].height = 22

    # Header row
    headers = [
        "Parent Supplier", "KPI ID", "KPI Name", "Max Score",
        "Scorecard Earned", "Independent Calc", "Delta",
        f"Tolerance (±{TOLERANCE})", "Status",
    ]
    for col, h in enumerate(headers, 1):
        hdr(ws1, 2, col, h)
    ws1.row_dimensions[2].height = 18

    # Data rows
    prev_parent = None
    row_num = 3
    for i, r in enumerate(detail_rows):
        is_alt = (i % 2 == 0)
        base_fill = ALT_FILL if is_alt else None
        parent = r["Parent Supplier"]
        status = r["Status"]

        if status == "PASS":
            status_fill = GREEN_FILL
        elif "MISSING" in status:
            status_fill = AMBER_FILL
        else:
            status_fill = RED_FILL

        parent_display = parent if parent != prev_parent else ""
        prev_parent = parent

        cell(ws1, row_num, 1, parent_display, fill=base_fill, bold=(parent_display != ""), align=LEFT)
        cell(ws1, row_num, 2, r["KPI ID"],    fill=base_fill, bold=True)
        cell(ws1, row_num, 3, r["KPI Name"],  fill=base_fill, align=LEFT)
        cell(ws1, row_num, 4, r["Max Score"], fill=base_fill)

        sc_val  = r["Scorecard Earned"]
        ind_val = r["Independent Calc Earned"]
        delta   = r["Delta"]

        sc_c = cell(ws1, row_num, 5, round(sc_val,  4) if sc_val  is not None else "—", fill=base_fill)
        in_c = cell(ws1, row_num, 6, round(ind_val, 4) if ind_val is not None else "—", fill=base_fill)
        de_c = cell(ws1, row_num, 7, round(delta,   6) if delta   is not None else "—", fill=base_fill)

        if sc_val is not None: sc_c.number_format = "0.0000"
        if ind_val is not None: in_c.number_format = "0.0000"
        if delta is not None:   de_c.number_format = "0.000000"

        cell(ws1, row_num, 8, TOLERANCE, fill=base_fill)
        cell(ws1, row_num, 9, status, fill=status_fill, bold=(status != "PASS"))

        ws1.row_dimensions[row_num].height = 15
        row_num += 1

    # Column widths
    col_widths = [44, 8, 28, 10, 17, 17, 12, 16, 22]
    for i, w in enumerate(col_widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # ════════════════════════════════════════════════════════════════════════
    # Sheet 2 — Summary
    # ════════════════════════════════════════════════════════════════════════
    ws2 = wb.create_sheet("Summary")
    ws2.freeze_panes = "A3"

    ws2.merge_cells("A1:F1")
    ws2["A1"].value = "Alignment Test Summary"
    ws2["A1"].font  = TITLE_FONT
    ws2["A1"].alignment = LEFT
    ws2.row_dimensions[1].height = 22

    # Per-KPI summary
    sum_headers = ["KPI ID", "KPI Name", "Checks Run", "PASS", "FAIL / MISSING", "Pass Rate %"]
    for col, h in enumerate(sum_headers, 1):
        hdr(ws2, 2, col, h)
    ws2.row_dimensions[2].height = 18

    kpi_stats: dict[str, dict] = {k["id"]: {"name": k["name"], "pass": 0, "fail": 0} for k in KPI_DEFS}
    for r in detail_rows:
        kid = r["KPI ID"]
        if r["Status"] == "PASS":
            kpi_stats[kid]["pass"] += 1
        else:
            kpi_stats[kid]["fail"] += 1

    for row_idx, kpi in enumerate(KPI_DEFS, 3):
        kid = kpi["id"]
        s = kpi_stats[kid]
        total_kpi = s["pass"] + s["fail"]
        pct = round(100.0 * s["pass"] / total_kpi, 1) if total_kpi else None
        fill = GREEN_FILL if (pct or 0) == 100 else (AMBER_FILL if (pct or 0) >= 80 else RED_FILL)

        cell(ws2, row_idx, 1, kid,         fill=fill, bold=True)
        cell(ws2, row_idx, 2, s["name"],   fill=fill, align=LEFT)
        cell(ws2, row_idx, 3, total_kpi,   fill=fill)
        cell(ws2, row_idx, 4, s["pass"],   fill=fill)
        cell(ws2, row_idx, 5, s["fail"],   fill=fill)
        cell(ws2, row_idx, 6, f"{pct}%" if pct is not None else "—", fill=fill, bold=True)
        ws2.row_dimensions[row_idx].height = 15

    # Overall totals row
    total_row = 3 + len(KPI_DEFS)
    total_checks = passed + failed
    overall_pct = round(100.0 * passed / total_checks, 1) if total_checks else 0
    fill_overall = GREEN_FILL if overall_pct == 100 else (AMBER_FILL if overall_pct >= 80 else RED_FILL)

    cell(ws2, total_row, 1, "TOTAL",        fill=fill_overall, bold=True)
    cell(ws2, total_row, 2, f"{len(sample)} parent suppliers tested", fill=fill_overall, bold=True, align=LEFT)
    cell(ws2, total_row, 3, total_checks,   fill=fill_overall, bold=True)
    cell(ws2, total_row, 4, passed,         fill=fill_overall, bold=True)
    cell(ws2, total_row, 5, failed,         fill=fill_overall, bold=True)
    cell(ws2, total_row, 6, f"{overall_pct}%", fill=fill_overall, bold=True)
    ws2.row_dimensions[total_row].height = 18

    # Metadata block
    meta_start = total_row + 2
    meta = [
        ("Test Mode",    mode),
        ("Suppliers Tested", len(sample)),
        ("Random Seed",  seed),
        ("Tolerance",    f"±{TOLERANCE} earned points"),
        ("Formula",      "softStretch  →  earned = max × attainment × (0.70 + 0.30 × percentile)"),
        ("Generated",    datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Overall Result", f"{'ALL PASS ✓' if failed == 0 else str(failed) + ' FAILURES ✗'}  ({overall_pct}%)"),
    ]
    for offset, (label, value) in enumerate(meta):
        ws2.cell(row=meta_start + offset, column=1, value=label).font  = BOLD_FONT
        ws2.cell(row=meta_start + offset, column=1).alignment          = LEFT
        ws2.merge_cells(
            start_row=meta_start + offset, start_column=2,
            end_row=meta_start + offset,   end_column=6
        )
        vc = ws2.cell(row=meta_start + offset, column=2, value=value)
        vc.font      = BODY_FONT
        vc.alignment = LEFT
        if label == "Overall Result":
            vc.fill = GREEN_FILL if failed == 0 else RED_FILL
            vc.font = BOLD_FONT

    # Column widths for summary sheet
    for col, w in enumerate([10, 32, 12, 10, 16, 14], 1):
        ws2.column_dimensions[get_column_letter(col)].width = w

    wb.save(out_path)
    print(f"\n  Excel report saved → {out_path}")



# ── Offline mode ────────────────────────────────────────────────────────────

def run_offline(n: int, seed: int, excel_path: Path | None = None) -> bool:
    """Load CSVs directly and compare against scorecard.py compute_scorecard()."""
    from scorecard import compute_scorecard  # type: ignore[import]
    from scorecard import compute_scorecard  # type: ignore[import]

    # 1. Load CSVs → cache dict
    cache: dict = {}
    cache_key_map = {
        "DOT": "dot_kpi", "SA": "supplier_assessment", "SC": "supplier_compliance",
        "PDIV": "price_divergence", "IC": "invoice_conformity", "IOT": "iot_kpi",
        "SM": "supplier_maturity", "ECL": "eclipse", "CO2": "co2_emission",
    }
    kpi_rows: dict[str, list[dict]] = {}
    for kpi in KPI_DEFS:
        path = DATA_DIR / kpi["csv"]
        if path.exists():
            df = pd.read_csv(path, low_memory=False)
            for col in df.select_dtypes(include=["object"]).columns:
                df[col] = df[col].astype(str).str.strip()
            rows = df.to_dict(orient="records")
        else:
            print(f"  [WARNING] Missing CSV: {kpi['csv']} — {kpi['name']} will be skipped")
            rows = []
        kpi_rows[kpi["id"]] = rows
        cache[cache_key_map[kpi["id"]]] = rows

    # 2. Run official scorecard computation
    result = compute_scorecard(cache, include_kpi_breakdown=True)

    # 3. Build scorecard lookup: parent → {kpi_id: earned}
    sc_lookup: dict[str, dict[str, float]] = {}
    for sc in result["scorecards"]:
        parent = sc["parentSupplier"]
        sc_lookup[parent] = {}
        for pillar in sc["pillars"]:
            for kpi_entry in pillar["kpis"]:
                if kpi_entry.get("applicable") and kpi_entry.get("earned") is not None:
                    sc_lookup[parent][kpi_entry["id"]] = kpi_entry["earned"]

    # 4. Independent computation from the same raw rows
    ind_earned: dict[str, dict[str, float]] = {}
    for kpi in KPI_DEFS:
        rows = kpi_rows[kpi["id"]]
        if not rows:
            continue
        ratios = _aggregate(kpi["id"], rows)
        ind_earned[kpi["id"]] = _softstretch_earned(
            ratios, kpi["direction"], kpi["floor"], kpi["target"], kpi["max"]
        )

    # 5. Random sample
    all_parents = list(sc_lookup.keys())
    random.seed(seed)
    sample = random.sample(all_parents, min(n, len(all_parents)))

    _print_header(len(sample), mode="Offline (CSV)")
    passed, failed, detail_rows = _compare_and_print(sample, sc_lookup, ind_earned)
    _print_summary(passed, failed)
    if excel_path:
        _export_excel(detail_rows, sample, passed, failed, "Offline (CSV)", seed, excel_path)
    return failed == 0


# ── API mode ────────────────────────────────────────────────────────────────

def run_api(api_base: str, n: int, seed: int, excel_path: Path | None = None) -> bool:
    """Fetch data from the live backend and compare scorecard vs raw KPI rollup."""
    import requests  # type: ignore[import]

    api_base = api_base.rstrip("/")

    # 1. Fetch all scorecard entries
    print(f"  Fetching scorecard from {api_base}/api/scorecard …")
    resp = requests.get(f"{api_base}/api/scorecard", timeout=30)
    resp.raise_for_status()
    scorecard_data = resp.json()

    all_parents = [sc["parentSupplier"] for sc in scorecard_data["scorecards"]]
    if not all_parents:
        print("  ERROR: Scorecard returned no parent suppliers.")
        return False

    # 2. Build scorecard lookup
    sc_lookup: dict[str, dict[str, float]] = {}
    for sc in scorecard_data["scorecards"]:
        parent = sc["parentSupplier"]
        sc_lookup[parent] = {}
        for pillar in sc.get("pillars", []):
            for kpi_entry in pillar.get("kpis", []):
                if kpi_entry.get("applicable") and kpi_entry.get("earned") is not None:
                    sc_lookup[parent][kpi_entry["id"]] = kpi_entry["earned"]

    # 3. Fetch raw KPI data from individual endpoints
    endpoint_map = {
        "DOT":  f"{api_base}/api/dot-kpi",
        "SA":   f"{api_base}/api/supplier-assessment",
        "SC":   f"{api_base}/api/supplier-compliance",
        "PDIV": f"{api_base}/api/price-divergence",
        "IC":   f"{api_base}/api/invoice-conformity",
        "IOT":  f"{api_base}/api/iot-kpi",
        "SM":   f"{api_base}/api/supplier-maturity",
        "ECL":  f"{api_base}/api/eclipse",
        "CO2":  f"{api_base}/api/co2-emission",
    }
    kpi_rows: dict[str, list[dict]] = {}
    for kid, url in endpoint_map.items():
        try:
            r = requests.get(url, timeout=30)
            r.raise_for_status()
            kpi_rows[kid] = r.json().get("data", [])
            print(f"  {kid:<6} → {len(kpi_rows[kid]):>5} rows  ({url})")
        except Exception as exc:
            print(f"  [WARNING] {kid}: failed to fetch {url} — {exc}")
            kpi_rows[kid] = []

    # 4. Independent computation
    ind_earned: dict[str, dict[str, float]] = {}
    for kpi in KPI_DEFS:
        rows = kpi_rows.get(kpi["id"], [])
        if not rows:
            continue
        ratios = _aggregate(kpi["id"], rows)
        ind_earned[kpi["id"]] = _softstretch_earned(
            ratios, kpi["direction"], kpi["floor"], kpi["target"], kpi["max"]
        )

    # 5. Random sample
    random.seed(seed)
    sample = random.sample(all_parents, min(n, len(all_parents)))

    print()
    _print_header(len(sample), mode=f"API ({api_base})")
    passed, failed, detail_rows = _compare_and_print(sample, sc_lookup, ind_earned)
    _print_summary(passed, failed)
    if excel_path:
        _export_excel(detail_rows, sample, passed, failed, f"API ({api_base})", seed, excel_path)
    return failed == 0


# ── Entry point ─────────────────────────────────────────────────────────────

if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Test alignment between Normalized Scorecard and individual KPI earned scores."
    )
    parser.add_argument(
        "--offline", action="store_true",
        help="Use CSV files directly (no server needed).",
    )
    parser.add_argument(
        "--api", default="http://127.0.0.1:8000",
        help="Base URL of the running backend (default: http://127.0.0.1:8000).",
    )
    parser.add_argument(
        "--n", type=int, default=20,
        help="Number of parent suppliers to sample (default: 20).",
    )
    parser.add_argument(
        "--seed", type=int, default=42,
        help="Random seed for reproducibility (default: 42).",
    )
    parser.add_argument(
        "--excel", action="store_true",
        help="Save results to an Excel report in backend/data/.",
    )
    args = parser.parse_args()

    from datetime import datetime
    excel_path: Path | None = None
    if args.excel:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        excel_path = DATA_DIR / f"scorecard_alignment_report_{ts}.xlsx"

    if args.offline:
        ok = run_offline(n=args.n, seed=args.seed, excel_path=excel_path)
    else:
        ok = run_api(api_base=args.api, n=args.n, seed=args.seed, excel_path=excel_path)

    sys.exit(0 if ok else 1)
